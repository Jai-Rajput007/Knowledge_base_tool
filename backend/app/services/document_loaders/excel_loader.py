"""Excel loader using openpyxl — formats sheets as labelled row text for LLM Q&A."""

from pathlib import Path
from typing import Union, Optional, List
import io

from app.services.document_loaders.base import (
    DocumentLoader, DocumentContent, ContentElement,
    ContentType, DocumentMetadata,
)
from app.core.logging import logger
from app.core.exceptions import DocumentProcessingError


class ExcelLoader(DocumentLoader):
    supported_extensions = [".xlsx", ".xls"]

    def load(self, file_path: Union[str, Path]) -> DocumentContent:
        file_path = Path(file_path)
        with open(file_path, "rb") as f:
            return self.load_from_bytes(f.read(), file_path.name)

    def load_from_bytes(self, content: bytes, filename: Optional[str] = None) -> DocumentContent:
        try:
            import openpyxl
        except ImportError:
            raise DocumentProcessingError("openpyxl not installed. Run: pip install openpyxl")

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except Exception:
            # .xls (old format) — try xlrd fallback
            return self._load_xls(content, filename)

        elements: List[ContentElement] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = [
                [self._cell_val(cell) for cell in row]
                for row in ws.iter_rows()
            ]
            # Drop completely empty rows
            rows = [r for r in rows if any(c != "" for c in r)]
            if not rows:
                continue

            # Sheet heading
            elements.append(ContentElement(
                type=ContentType.HEADING,
                content=f"Sheet: {sheet_name}",
                level=1,
            ))

            for chunk in self._format_sheet_chunks(rows, sheet_name):
                elements.append(ContentElement(
                    type=ContentType.TABLE,
                    content=chunk,
                    metadata={"sheet": sheet_name, "source": "openpyxl"},
                ))

        metadata = DocumentMetadata(
            title=filename,
            source_file=filename,
            file_size=len(content),
            page_count=len(wb.sheetnames),
        )
        doc = DocumentContent(elements=elements, metadata=metadata)
        metadata.word_count = len(doc.to_text().split())
        logger.info(f"Excel loaded: {filename} — {len(wb.sheetnames)} sheet(s), {len(elements)} elements")
        return doc

    def _cell_val(self, cell) -> str:
        if cell.value is None:
            return ""
        # Collapse internal newlines so multiline cells don't break pipe-delimited format
        return " ".join(str(cell.value).split())

    def _format_sheet_chunks(self, rows: List[List[str]], sheet_name: str, max_rows: int = 5) -> List[str]:
        """Split a sheet into chunks of max_rows data rows each."""
        if not rows:
            return []

        def is_header(row: List[str]) -> bool:
            import re
            non_empty = [c for c in row if c]
            if not non_empty:
                return False
            numeric = sum(1 for c in non_empty if re.match(r'^[\d,.\-%()\s]+$', c))
            return numeric < len(non_empty) / 2

        has_header = len(rows) > 1 and is_header(rows[0])
        headers = rows[0] if has_header else []
        data_rows = rows[1:] if has_header else rows

        total_data_rows = len(data_rows)
        chunks = []
        for i in range(0, max(total_data_rows, 1), max_rows):
            batch = data_rows[i:i + max_rows]
            part = i // max_rows + 1
            total = (total_data_rows - 1) // max_rows + 1 if total_data_rows > 0 else 1
            row_range = f"rows {i+1}-{min(i+max_rows, total_data_rows)}"
            if total > 1:
                label = f"[Sheet: {sheet_name} | Total rows: {total_data_rows} | Part {part}/{total} ({row_range})]"
            else:
                label = f"[Sheet: {sheet_name} | Total rows: {total_data_rows}]"
            lines = [label]
            # Characters used as Gantt/bar chart fill — skip these columns entirely
            _CHART_CHARS = {"█", "■", "▓", "░", "▪", "●"}

            if has_header and headers:
                for row in batch:
                    parts = []
                    for j, val in enumerate(row):
                        if not val:
                            continue
                        # Skip Gantt chart bar values
                        if val in _CHART_CHARS or all(c in _CHART_CHARS for c in val):
                            continue
                        lbl = headers[j] if j < len(headers) and headers[j] else f"Col{j+1}"
                        parts.append(f"{lbl}: {val}")
                    if parts:
                        lines.append(" | ".join(parts))
            else:
                for row in batch:
                    line = " | ".join(c for c in row if c and c not in _CHART_CHARS)
                    if line:
                        lines.append(line)
            if len(lines) > 1:
                chunks.append("\n".join(lines))

        return chunks if chunks else []

    def _format_sheet(self, rows: List[List[str]], sheet_name: str) -> str:
        """
        Format a sheet as labelled prose so the LLM can answer questions accurately.

        If the first row looks like headers, each subsequent row becomes:
            Header1: Value | Header2: Value | ...

        Otherwise rows are joined with pipe separators.
        """
        if not rows:
            return ""

        def is_header(row: List[str]) -> bool:
            non_empty = [c for c in row if c]
            if not non_empty:
                return False
            import re
            numeric = sum(1 for c in non_empty if re.match(r'^[\d,.\-%()\s]+$', c))
            return numeric < len(non_empty) / 2

        lines = [f"[Sheet: {sheet_name}]"]

        if len(rows) > 1 and is_header(rows[0]):
            headers = rows[0]
            for row in rows[1:]:
                parts = []
                for i, val in enumerate(row):
                    if not val:
                        continue
                    label = headers[i] if i < len(headers) and headers[i] else f"Col{i+1}"
                    parts.append(f"{label}: {val}")
                if parts:
                    lines.append(" | ".join(parts))
        else:
            for row in rows:
                line = " | ".join(c for c in row if c)
                if line:
                    lines.append(line)

        return "\n".join(lines)

    def _load_xls(self, content: bytes, filename: Optional[str]) -> DocumentContent:
        """Fallback for legacy .xls files using xlrd."""
        try:
            import xlrd
        except ImportError:
            raise DocumentProcessingError(
                ".xls format requires xlrd. Run: pip install xlrd"
            )
        try:
            wb = xlrd.open_workbook(file_contents=content)
        except Exception as e:
            raise DocumentProcessingError(f"Failed to read .xls file: {e}")

        elements: List[ContentElement] = []
        for sheet in wb.sheets():
            rows = [
                [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
                for r in range(sheet.nrows)
            ]
            rows = [r for r in rows if any(c for c in r)]
            if not rows:
                continue
            elements.append(ContentElement(
                type=ContentType.HEADING, content=f"Sheet: {sheet.name}", level=1
            ))
            sheet_text = self._format_sheet(rows, sheet.name)
            if sheet_text:
                elements.append(ContentElement(
                    type=ContentType.TABLE,
                    content=sheet_text,
                    metadata={"sheet": sheet.name, "source": "xlrd"},
                ))

        metadata = DocumentMetadata(title=filename, source_file=filename, file_size=len(content))
        doc = DocumentContent(elements=elements, metadata=metadata)
        metadata.word_count = len(doc.to_text().split())
        return doc
